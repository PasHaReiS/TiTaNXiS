"""Backend tests for iteration 13: PATCH /scores, Excel export Sheet 3 (Sıralama)."""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/") or "https://oyun-loncasi.preview.emergentagent.com"


@pytest.fixture(scope="module")
def admin_token():
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"username": "admin", "password": "admin123"}, timeout=30)
    assert r.status_code == 200, r.text
    return r.json()["token"]


@pytest.fixture(scope="module")
def auth_headers(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


class TestExportSralama:
    def test_export_xlsx_sheet3_ranks(self, auth_headers):
        r = requests.get(f"{BASE_URL}/api/export/xlsx", headers=auth_headers, timeout=60)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert "Sıralama Listesi" in wb.sheetnames
        ws = wb["Sıralama Listesi"]
        headers = [c.value for c in ws[1]]
        assert headers == ["Sıra", "Üye", "Rütbe", "İttifak", "Puan"], headers
        # column A rows 2..N are 1..N-1
        data_rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[1]]  # has name
        assert len(data_rows) > 0
        for i, row in enumerate(data_rows, start=1):
            assert row[0] == i, f"row {i}: expected {i}, got {row[0]}"

        # Total should match /api/members count
        members = requests.get(f"{BASE_URL}/api/members", headers=auth_headers, timeout=30).json()
        assert len(data_rows) == len(members), f"sheet3={len(data_rows)} vs members={len(members)}"


class TestPatchScore:
    def test_patch_score_updates(self, auth_headers):
        # find any existing score
        scores = requests.get(f"{BASE_URL}/api/scores?limit=1", headers=auth_headers, timeout=30).json()
        if not scores:
            pytest.skip("no scores exist to patch")
        s = scores[0]
        original_points = s["points"]
        new_points = original_points + 7
        r = requests.patch(
            f"{BASE_URL}/api/scores/{s['id']}",
            headers=auth_headers,
            json={"points": new_points, "multiplier": s.get("multiplier", 1), "event_id": s["event_id"], "note": s.get("note")},
            timeout=30,
        )
        assert r.status_code in (200, 204), r.text
        # verify persisted
        after = requests.get(f"{BASE_URL}/api/scores?limit=2000", headers=auth_headers, timeout=30).json()
        match = next((x for x in after if x["id"] == s["id"]), None)
        assert match is not None
        assert match["points"] == new_points
        # restore
        requests.patch(
            f"{BASE_URL}/api/scores/{s['id']}",
            headers=auth_headers,
            json={"points": original_points, "multiplier": s.get("multiplier", 1), "event_id": s["event_id"], "note": s.get("note")},
            timeout=30,
        )
